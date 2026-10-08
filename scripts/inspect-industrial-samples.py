"""Inspect industrial examples with bounded memory; never change source data."""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import pandas as pd


PROJECT = Path(__file__).resolve().parents[1]
SAMPLES = PROJECT.parent / "AutoDecision_Sample"
DATASETS = ("车流预测", "城市配送智能", "化学工厂", "steel+industry+energy+consumption",
            "ai4i+2020+predictive+maintenance+dataset")


def text_document(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        return "\n".join("".join(p.itertext()) for p in root.findall(".//w:p", ns))
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    return path.read_text(encoding="utf-8-sig", errors="replace")


def csv_profile(path: Path) -> dict:
    result = {}
    encodings = ("utf-8-sig", "cp1252", "gb18030") if "Data_for_" in str(path) else ("utf-8-sig", "gb18030", "cp1252")
    for encoding in encodings:
        try:
            with path.open(encoding=encoding) as stream:
                preview = stream.read(8192)
            try:
                delimiter = csv.Sniffer().sniff(preview, delimiters=",;\t").delimiter
            except csv.Error:
                header = preview.splitlines()[0]
                delimiter = max((",", ";", "\t"), key=header.count)
            result = {"encoding": encoding, "delimiter": delimiter, "preview": preview[:1800]}
            total = 0
            nulls: dict[str, int] = {}
            for index, chunk in enumerate(pd.read_csv(path, sep=delimiter, encoding=encoding, chunksize=20000, low_memory=False)):
                if index == 0:
                    result["columns"] = list(chunk.columns)
                    result["sample"] = chunk.head(3).where(chunk.head(3).notna(), None).to_dict("records")
                    result["sample_dtypes"] = {str(k): str(v) for k, v in chunk.dtypes.items()}
                total += len(chunk)
                for key, value in chunk.isna().sum().items():
                    nulls[str(key)] = nulls.get(str(key), 0) + int(value)
            result.update(rows=total, null_counts=nulls)
            return result
        except UnicodeDecodeError:
            continue
        except Exception as exc:
            result["table_parse_error"] = str(exc)
            return result
    return result


def workbook_profile(path: Path) -> dict:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = []
            for sheet in book:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    if any(value is not None for value in row):
                        rows.append(list(row))
                    if len(rows) >= 6:
                        break
                sheets.append({"sheet": sheet.title, "used_rows": sheet.max_row,
                               "used_columns": sheet.max_column, "first_nonempty_rows": rows})
            return {"sheets": sheets}
        finally:
            book.close()
    book = pd.ExcelFile(path)
    return {"sheets": [{"sheet": sheet, "first_rows": pd.read_excel(book, sheet_name=sheet, header=None, nrows=8).fillna("").values.tolist()}
                       for sheet in book.sheet_names]}


def summarize_usage(root: Path) -> dict:
    stages = []
    for path in sorted(root.rglob("llm_usage.jsonl")):
        rows = [json.loads(line) for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]
        names: dict[str, dict] = {}
        for row in rows:
            if str(row.get("source", "")) in {"local_cache", "cache"} or row.get("cache_hit_local"):
                continue
            name = str(row.get("prompt_name") or row.get("component") or "unknown")
            item = names.setdefault(name, {"calls": 0, "input_tokens": 0, "output_tokens": 0, "cache_known_calls": 0, "cached_tokens": 0})
            item["calls"] += 1
            item["input_tokens"] += int(row.get("prompt_tokens") or 0)
            item["output_tokens"] += int(row.get("completion_tokens") or 0)
            item["cache_known_calls"] += int(bool(row.get("provider_cache_tokens_known")))
            item["cached_tokens"] += int(row.get("prompt_cache_hit_tokens") or 0)
        stages.append({"path": str(path.relative_to(root)), "prompts": dict(sorted(names.items(), key=lambda item: -item[1]["input_tokens"]))})
    return {"stages": stages}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=PROJECT / "runs" / "industrial-examples-20260907" / "inspection")
    parser.add_argument("--dataset", action="append")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    roots = [SAMPLES / name for name in DATASETS]
    roots += sorted(SAMPLES.glob("Data_for_*"))
    roots = [path for path in roots if path.is_dir() and (not args.dataset or path.name in args.dataset)]
    for root in roots:
        inventory = []
        for path in sorted(root.rglob("*")):
            if not path.is_file():
                continue
            row = {"file": str(path.relative_to(root)), "bytes": path.stat().st_size}
            try:
                suffix = path.suffix.lower()
                if suffix == ".csv":
                    row.update(csv_profile(path))
                elif suffix in {".xls", ".xlsx"}:
                    row.update(workbook_profile(path))
                elif suffix in {".txt", ".md", ".docx", ".pdf"}:
                    row["text"] = text_document(path)
            except Exception as exc:
                row["error"] = str(exc)
            inventory.append(row)
        output = args.output / (root.name[:45] + ".json")
        output.write_text(json.dumps({"dataset": root.name, "root": str(root), "files": inventory}, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        print(json.dumps({"dataset": root.name, "files": len(inventory), "output": str(output), "errors": [row["file"] for row in inventory if "error" in row]}, ensure_ascii=False), flush=True)
    prior = PROJECT / "runs" / "Live validation - store sales"
    (args.output / "prior-usage.json").write_text(json.dumps(summarize_usage(prior), ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
